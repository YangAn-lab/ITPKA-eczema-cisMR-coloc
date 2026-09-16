from docx import Document
d = Document('/mnt/results/P2-02_manuscript_Genomics_submission.docx')
print('paragraphs:', len(d.paragraphs), '| tables:', len(d.tables))
print('title:', d.paragraphs[0].text[:80])
print('last ref:', [p.text for p in d.paragraphs if p.text.strip().startswith('51.')][0][:70])