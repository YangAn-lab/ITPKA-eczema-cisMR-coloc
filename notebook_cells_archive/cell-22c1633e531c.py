from docx import Document
d = Document('/workspace/P2-02_manuscript_Genomics_submission.docx')

# Table dimensions
for k, t in enumerate(d.tables, 1):
    print(f"Table {k}: {len(t.rows)} rows x {len(t.columns)} cols | header: {[c.text[:22] for c in t.rows[0].cells][:4]}")

# Reference count: paragraphs starting with a number+period pattern in the References block
import re
texts = [p.text for p in d.paragraphs]
ref_start = next(i for i, t in enumerate(texts) if t.strip() == 'References')
refs = [t for t in texts[ref_start+1:] if re.match(r'^\d+\.\s', t.strip())]
print('\nReferences found:', len(refs), '| first:', refs[0][:60], '| last:', refs[-1][:60])

# Spot-check key sections present
for kw in ['Abstract', 'Highlights', 'Declarations', 'Data availability', 'Figure legends']:
    print(kw, '->', any(kw.lower() in t.lower() for t in texts))
print('\nTotal chars in docx paragraphs:', sum(len(t) for t in texts))