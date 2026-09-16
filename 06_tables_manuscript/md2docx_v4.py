"""Convert the Genomics submission markdown to a single-column editable .docx.
Handles: #/##/### headings, **bold** inline, pipe tables, --- rules, paragraphs.
Output: /workspace/P2-02_manuscript_v4.docx
"""
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_LINE_SPACING

SRC = "/workspace/P2-02_manuscript_v4.md"
DST = "/workspace/P2-02_manuscript_v4.docx"

doc = Document()
# base style
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
    try:
        doc.styles[s].font.name = "Times New Roman"
    except Exception:
        pass

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def add_par(text, style_name=None):
    p = doc.add_paragraph(style=style_name)
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        p.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])
    return p


def flush_table(rows):
    # rows: list of list-of-cells; rows[1] is the |---| separator
    header = rows[0]
    body = [r for r in rows[1:] if not all(set(c.strip()) <= set("-: ") for c in r)]
    t = doc.add_table(rows=len(body) + 1, cols=len(header))
    t.style = "Table Grid"
    for j, c in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(c.strip())
        run.bold = True
        run.font.size = Pt(9)
    for i, r in enumerate(body, start=1):
        for j in range(len(header)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            txt = r[j].strip() if j < len(r) else ""
            m = BOLD_RE.fullmatch(txt)
            run = cell.paragraphs[0].add_run(BOLD_RE.sub(r"\1", txt))
            run.font.size = Pt(9)
            if m:
                run.bold = True
    doc.add_paragraph()


lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1
        continue
    if line.strip() == "---":
        i += 1
        continue
    if line.startswith("|"):
        tbl = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c for c in lines[i].strip().strip("|").split("|")]
            tbl.append(cells)
            i += 1
        flush_table(tbl)
        continue
    if line.startswith("### "):
        add_par(line[4:], "Heading 3")
    elif line.startswith("## "):
        add_par(line[3:], "Heading 2")
    elif line.startswith("# "):
        add_par(line[2:], "Heading 1")
    else:
        add_par(line)
    i += 1

doc.save(DST)
print("saved", DST)

# read-back verification
from docx import Document as D2
d = D2(DST)
print("paragraphs:", len(d.paragraphs), "| tables:", len(d.tables))
print("first para:", d.paragraphs[0].text[:90])
